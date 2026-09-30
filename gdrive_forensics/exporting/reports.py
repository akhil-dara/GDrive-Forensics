"""Report rows and writers. Column order of v1 reports is preserved; new columns are appended."""
from __future__ import annotations

import csv
import json
import re
from typing import Optional

from ..core.formatting import format_size, format_with_offset, now_iso
from ..core.mime import FOLDER_MIME

EXPORT_REPORT_COLUMNS = [
    "File_ID", "Name", "MIME_Type", "Item_Type", "Size_Bytes", "Size_Readable", "Drive_Path", "Local_Path",
    "Owner_Name", "Owner_Email", "Starred", "Trashed", "Public", "Shortcut", "MD5", "SHA1", "SHA256",
    "Web_View_Link", "Shared_With_JSON",
    # appended in v2
    "Downloaded", "Local_MD5", "Local_SHA256", "Hash_Verified", "Hash_Algorithm", "Error",
    "Created_Time", "Modified_Time", "Created_Time_Local", "Modified_Time_Local",
]

METADATA_REPORT_COLUMNS = [
    "File_ID", "Name", "MIME_Type", "Item_Type", "Size_Bytes", "Size_Readable", "Drive_Path", "Local_Path",
    "Owner_Name", "Owner_Email", "Starred", "Trashed", "Is_Public", "Is_Shortcut", "MD5", "SHA1", "SHA256",
    "Web_Link", "Shared_With", "Created_Time", "Modified_Time", "Created_Time_Local", "Modified_Time_Local",
    "Timezone",
]


def _yes_no(value) -> str:
    return "Yes" if value else "No"


def _verified(value) -> str:
    return {True: "Yes", False: "No"}.get(value, "N/A")


def build_report_entry(record: dict, *, drive_path: str, permissions: list, local_path: Optional[str],
                       tz_name: str, download=None, error: Optional[str] = None) -> dict:
    size = record.get("size") or 0
    created, modified = record.get("created_time"), record.get("modified_time")
    return {
        "file_id": record["id"], "name": record.get("name"), "mime_type": record.get("mime_type"),
        "size_bytes": size, "size_readable": format_size(size),
        "item_type": "folder" if record.get("mime_type") == FOLDER_MIME else "file",
        "drive_path": drive_path, "local_path": local_path or "Not downloaded",
        "owner_name": record.get("owner_name"), "owner_email": record.get("owner_email"),
        "starred": bool(record.get("starred")), "trashed": bool(record.get("trashed")),
        "is_public": bool(record.get("is_public")), "is_shortcut": bool(record.get("is_shortcut")),
        "md5_checksum": record.get("md5_checksum"), "sha1_checksum": record.get("sha1_checksum"),
        "sha256_checksum": record.get("sha256_checksum"), "web_view_link": record.get("web_view_link"),
        "shared_with": permissions, "created_time": created, "modified_time": modified,
        "created_time_local": format_with_offset(created, tz_name),
        "modified_time_local": format_with_offset(modified, tz_name),
        "downloaded": download is not None,
        "local_md5": download.md5 if download else None,
        "local_sha256": download.sha256 if download else None,
        "hash_algorithm": download.algorithm if download else None,
        "hash_verified": download.verified if download else None,
        "error": error,
    }


def _export_row(e: dict) -> list:
    return [e["file_id"], e["name"], e["mime_type"], e["item_type"], e["size_bytes"], e["size_readable"],
            e["drive_path"], e["local_path"], e["owner_name"], e["owner_email"], _yes_no(e["starred"]),
            _yes_no(e["trashed"]), _yes_no(e["is_public"]), _yes_no(e["is_shortcut"]), e["md5_checksum"] or "",
            e["sha1_checksum"] or "", e["sha256_checksum"] or "", e["web_view_link"] or "",
            json.dumps(e["shared_with"], ensure_ascii=False), _yes_no(e["downloaded"]), e["local_md5"] or "",
            e["local_sha256"] or "", _verified(e["hash_verified"]), e["hash_algorithm"] or "", e["error"] or "",
            e["created_time"] or "", e["modified_time"] or "", e["created_time_local"] or "",
            e["modified_time_local"] or ""]


def _metadata_row(e: dict, tz_display: str) -> list:
    return [e["file_id"], e["name"], e["mime_type"], e["item_type"], e["size_bytes"], e["size_readable"],
            e["drive_path"], e["local_path"], e["owner_name"], e["owner_email"], _yes_no(e["starred"]),
            _yes_no(e["trashed"]), _yes_no(e["is_public"]), _yes_no(e["is_shortcut"]), e["md5_checksum"] or "",
            e["sha1_checksum"] or "", e["sha256_checksum"] or "", e["web_view_link"] or "",
            json.dumps(e["shared_with"], ensure_ascii=False), e["created_time"] or "", e["modified_time"] or "",
            e["created_time_local"] or "", e["modified_time_local"] or "", tz_display]


def write_export_csv(path, entries: list[dict]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(EXPORT_REPORT_COLUMNS)
        writer.writerows(_export_row(e) for e in entries)


def write_export_json(path, entries: list[dict], metadata: dict) -> None:
    payload = {"metadata": {"generated_at": now_iso(), **metadata}, "files": entries}
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, indent=2, ensure_ascii=False, default=str)


def write_metadata_csv(path, entries: list[dict], tz_display: str) -> None:
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(METADATA_REPORT_COLUMNS)
        writer.writerows(_metadata_row(e, tz_display) for e in entries)


def write_metadata_json(path, entries: list[dict], tz_display: str, filters_summary: str) -> None:
    payload = {"generated_at": now_iso(), "timezone": tz_display, "total_records": len(entries),
               "filters_applied": filters_summary, "records": entries}
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, indent=2, ensure_ascii=False, default=str)


# Characters XML 1.0 (and so openpyxl) refuses in cell text; tab, LF and CR are legal and kept.
_XLSX_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def xlsx_text(value: str) -> str:
    """Render control characters visibly (\\x07 -> the 4-char text `\\x07`) instead of aborting the workbook.

    CSV and JSON reports keep the exact names; this only affects the XLSX rendering.
    """
    return _XLSX_ILLEGAL.sub(lambda m: f"\\x{ord(m.group()):02x}", value)


def write_metadata_xlsx(path, entries: list[dict], tz_display: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Metadata Report"
    ws.append(METADATA_REPORT_COLUMNS)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for cell in ws[1]:
        cell.font, cell.fill = header_font, header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    size_index = METADATA_REPORT_COLUMNS.index("Size_Bytes")
    for e in entries:
        row = _metadata_row(e, tz_display)
        ws.append([int(v or 0) if i == size_index else ("" if v is None else xlsx_text(str(v)))
                   for i, v in enumerate(row)])
        # Drive names/paths/owners are untrusted: text starting with "=" must stay the exact text, never a
        # live formula (openpyxl would store it as data_type "f").
        for cell in ws[ws.max_row]:
            if cell.data_type == "f":
                cell.data_type = "s"
    for col_index, header in enumerate(METADATA_REPORT_COLUMNS, start=1):
        sample = [len(str(ws.cell(row=r, column=col_index).value or "")) for r in range(2, min(ws.max_row, 101) + 1)]
        ws.column_dimensions[get_column_letter(col_index)].width = min(max([len(header), *sample]) + 2, 50)
    ws.freeze_panes = "A2"
    wb.save(path)
