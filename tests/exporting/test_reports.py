import csv
import json

from openpyxl import load_workbook

from gdrive_forensics.drive.downloads import DownloadResult
from gdrive_forensics.exporting import reports

REC = {"id": "f1", "name": "a.pdf", "mime_type": "application/pdf", "size": 2048, "owner_name": "Me",
       "owner_email": "me@x.com", "starred": 1, "trashed": 0, "is_public": 1, "is_shortcut": 0,
       "md5_checksum": "m", "sha1_checksum": None, "sha256_checksum": None, "web_view_link": "L",
       "created_time": "2025-01-01T00:00:00Z", "modified_time": "2025-01-02T00:00:00Z"}
PERMS = [{"name": "Bob", "email": "bob@x.com", "role": "writer", "type": "user"}]


def entry(download=None, local_path=None):
    return reports.build_report_entry(REC, drive_path="Case / a.pdf", permissions=PERMS, local_path=local_path,
                                      tz_name="Asia/Kolkata", download=download)


def test_build_report_entry():
    e = entry()
    assert e["local_path"] == "Not downloaded" and e["size_readable"] == "2.00 KB" and e["item_type"] == "file"
    assert e["created_time_local"] == "2025-01-01 05:30:00 (UTC+05:30)" and e["shared_with"] == PERMS
    d = DownloadResult(path="/x/a.pdf", size=2048, md5="m", sha1="s1", sha256="s2", algorithm="md5",
                       expected_hash="m", verified=True)
    e = entry(download=d, local_path="Export/a.pdf")
    assert e["downloaded"] and e["hash_verified"] is True and e["local_sha256"] == "s2"


def test_export_csv_and_json(tmp_path):
    rows = [entry()]
    reports.write_export_csv(tmp_path / "r.csv", rows)
    with open(tmp_path / "r.csv", encoding="utf-8") as fh:
        data = list(csv.reader(fh))
    assert data[0] == reports.EXPORT_REPORT_COLUMNS
    assert data[0][:19] == ["File_ID", "Name", "MIME_Type", "Item_Type", "Size_Bytes", "Size_Readable", "Drive_Path",
                            "Local_Path", "Owner_Name", "Owner_Email", "Starred", "Trashed", "Public", "Shortcut",
                            "MD5", "SHA1", "SHA256", "Web_View_Link", "Shared_With_JSON"]
    assert data[1][10] == "Yes" and json.loads(data[1][18]) == PERMS and data[1][22] == "N/A"
    reports.write_export_json(tmp_path / "r.json", rows, {"total_records": 1})
    payload = json.loads((tmp_path / "r.json").read_text(encoding="utf-8"))
    assert payload["metadata"]["total_records"] == 1 and payload["files"][0]["file_id"] == "f1"


def test_metadata_reports(tmp_path):
    rows = [entry()]
    reports.write_metadata_csv(tmp_path / "m.csv", rows, "UTC (UTC+00:00)")
    with open(tmp_path / "m.csv", encoding="utf-8") as fh:
        data = list(csv.reader(fh))
    assert data[0] == reports.METADATA_REPORT_COLUMNS and len(data[0]) == 24 and data[1][-1] == "UTC (UTC+00:00)"
    reports.write_metadata_json(tmp_path / "m.json", rows, "UTC (UTC+00:00)", "No filters active")
    payload = json.loads((tmp_path / "m.json").read_text(encoding="utf-8"))
    assert payload["filters_applied"] == "No filters active" and payload["records"][0]["name"] == "a.pdf"
    reports.write_metadata_xlsx(tmp_path / "m.xlsx", rows, "UTC (UTC+00:00)")
    ws = load_workbook(tmp_path / "m.xlsx").active
    assert ws.title == "Metadata Report" and ws.freeze_panes == "A2"
    assert [c.value for c in ws[1]] == reports.METADATA_REPORT_COLUMNS
    assert ws.cell(row=2, column=5).value == 2048 and ws.cell(row=1, column=1).font.bold


def test_metadata_xlsx_stores_formula_like_text_as_values(tmp_path):
    evil = '=HYPERLINK("http://evil.example/","invoice.pdf")'
    rec = dict(REC, name=evil, owner_name="=1+1")
    e = reports.build_report_entry(rec, drive_path="Case / " + evil, permissions=PERMS, local_path=None,
                                   tz_name="UTC")
    reports.write_metadata_xlsx(tmp_path / "m.xlsx", [e], "UTC (UTC+00:00)")
    ws = load_workbook(tmp_path / "m.xlsx").active
    name_cell = ws.cell(row=2, column=2)
    assert name_cell.data_type == "s" and name_cell.value == evil  # exact text preserved, not a live formula
    owner_cell = ws.cell(row=2, column=9)
    assert owner_cell.data_type == "s" and owner_cell.value == "=1+1"
    assert ws.cell(row=2, column=7).value == "Case / " + evil
    assert all(c.data_type != "f" for row in ws.iter_rows() for c in row)


def test_metadata_xlsx_escapes_control_characters_visibly(tmp_path):
    rec = dict(REC, name="bell\x07.pdf", owner_name="Esc\x1b[0m Null\x00")
    e = reports.build_report_entry(rec, drive_path="Case / bell\x07.pdf", permissions=PERMS, local_path=None,
                                   tz_name="UTC")
    e["mime_type"] = "tab\there\nline"                       # XML-legal whitespace stays verbatim
    reports.write_metadata_xlsx(tmp_path / "m.xlsx", [e], "UTC (UTC+00:00)")   # must not abort
    ws = load_workbook(tmp_path / "m.xlsx").active
    assert ws.cell(row=2, column=2).value == r"bell\x07.pdf"           # the 4-char text \x07
    assert ws.cell(row=2, column=7).value == r"Case / bell\x07.pdf"
    assert ws.cell(row=2, column=9).value == r"Esc\x1b[0m Null\x00"
    assert ws.cell(row=2, column=3).value == "tab\there\nline"
